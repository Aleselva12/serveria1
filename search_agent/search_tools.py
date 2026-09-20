import json
import os
from pathlib import Path
from typing import Iterable

from dotenv import load_dotenv
from docx import Document
from langchain_core.tools import tool


load_dotenv()

DEFAULT_ROOT = Path(__file__).resolve().parents[1]
KNOWLEDGE_ROOT = Path(
    os.getenv("CORA_KNOWLEDGE_ROOT", str(DEFAULT_ROOT))
).expanduser().resolve()

BLOCKED_PARTS = {".git", ".venv", "__pycache__", "node_modules"}
BLOCKED_NAMES = {".env", "credentials.json", "token.json"}

TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".csv", ".py",
    ".yaml", ".yml", ".toml", ".html", ".css", ".js",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".docx"}
MAX_FILE_BYTES = int(os.getenv("CORA_MAX_DOCUMENT_BYTES", "5000000"))
MAX_TEXT_CHARS = int(os.getenv("CORA_MAX_DOCUMENT_CHARS", "120000"))


def _safe_path(relative_path: str) -> Path:
    candidate = (KNOWLEDGE_ROOT / relative_path).resolve()

    if candidate != KNOWLEDGE_ROOT and KNOWLEDGE_ROOT not in candidate.parents:
        raise ValueError("Accesso esterno alla cartella documenti vietato.")

    relative = candidate.relative_to(KNOWLEDGE_ROOT)
    if any(part in BLOCKED_PARTS for part in relative.parts):
        raise ValueError("Cartella protetta.")

    if candidate.name.lower() in BLOCKED_NAMES:
        raise ValueError("File protetto.")

    return candidate


def _iter_documents(directory: str = ".", recursive: bool = True) -> Iterable[Path]:
    base = _safe_path(directory)
    if not base.exists() or not base.is_dir():
        return []

    iterator = base.rglob("*") if recursive else base.glob("*")
    return (
        path for path in iterator
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
        and not any(part in BLOCKED_PARTS for part in path.relative_to(KNOWLEDGE_ROOT).parts)
        and path.name.lower() not in BLOCKED_NAMES
    )


def _read_docx(path: Path) -> str:
    document = Document(path)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        parts.append(f"\n[TABELLA {table_index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def _read_document(path: Path) -> str:
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError(
            f"File troppo grande: limite {MAX_FILE_BYTES} byte."
        )

    if path.suffix.lower() == ".docx":
        text = _read_docx(path)
    elif path.suffix.lower() in TEXT_EXTENSIONS:
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError("Formato non supportato.")

    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n\n[TESTO TRONCATO PER LIMITE DI SICUREZZA]"

    return text


@tool
def list_local_documents(
    directory: str = ".",
    extension: str = "",
    recursive: bool = True,
) -> str:
    """Elenca i documenti locali accessibili al Local Research Agent."""
    try:
        normalized_extension = extension.strip().lower()
        if normalized_extension and not normalized_extension.startswith("."):
            normalized_extension = "." + normalized_extension

        results = []
        for path in _iter_documents(directory, recursive):
            if normalized_extension and path.suffix.lower() != normalized_extension:
                continue

            stat = path.stat()
            results.append({
                "path": str(path.relative_to(KNOWLEDGE_ROOT)),
                "type": path.suffix.lower(),
                "size_bytes": stat.st_size,
                "modified_timestamp": stat.st_mtime,
            })

            if len(results) >= 200:
                break

        return json.dumps({
            "knowledge_root": str(KNOWLEDGE_ROOT),
            "count": len(results),
            "documents": results,
        }, ensure_ascii=False, indent=2)
    except Exception as error:
        return f"Errore durante l'elenco dei documenti: {error}"


@tool
def read_local_document(relative_path: str) -> str:
    """Legge il contenuto di un documento locale supportato, inclusi file Word .docx."""
    try:
        path = _safe_path(relative_path)

        if not path.exists() or not path.is_file():
            return "Il documento richiesto non esiste."

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            return (
                "Formato non supportato. Formati disponibili: "
                + ", ".join(sorted(SUPPORTED_EXTENSIONS))
            )

        text = _read_document(path)
        return (
            f"SOURCE_PATH: {path.relative_to(KNOWLEDGE_ROOT)}\n"
            f"SOURCE_TYPE: {path.suffix.lower()}\n"
            f"--- DOCUMENT CONTENT ---\n{text}"
        )
    except Exception as error:
        return f"Errore durante la lettura del documento: {error}"


@tool
def search_local_documents(
    query: str,
    directory: str = ".",
    max_results: int = 10,
) -> str:
    """
    Cerca parole o frasi nei documenti locali e restituisce estratti con il percorso
    della fonte. È una ricerca lessicale locale, non usa Internet.
    """
    try:
        terms = [term.casefold() for term in query.split() if len(term.strip()) >= 2]
        if not terms:
            return "La query non contiene termini utili."

        matches = []
        for path in _iter_documents(directory, recursive=True):
            try:
                text = _read_document(path)
            except Exception:
                continue

            folded = text.casefold()
            score = sum(folded.count(term) for term in terms)
            if score <= 0:
                continue

            positions = [folded.find(term) for term in terms if folded.find(term) >= 0]
            first_position = min(positions) if positions else 0
            start = max(0, first_position - 350)
            end = min(len(text), first_position + 850)
            excerpt = text[start:end].strip()

            matches.append({
                "path": str(path.relative_to(KNOWLEDGE_ROOT)),
                "score": score,
                "excerpt": excerpt,
            })

        matches.sort(key=lambda item: item["score"], reverse=True)
        matches = matches[:max(1, min(max_results, 30))]

        if not matches:
            return "Nessun documento locale contiene i termini cercati."

        return json.dumps({
            "query": query,
            "results": matches,
        }, ensure_ascii=False, indent=2)
    except Exception as error:
        return f"Errore durante la ricerca locale: {error}"


LOCAL_RESEARCH_TOOLS = [
    list_local_documents,
    search_local_documents,
    read_local_document,
]
