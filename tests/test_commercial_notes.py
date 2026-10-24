"""Isolated commercial Word workflow tests; no real client documents or DB."""
import tempfile
import unittest
from pathlib import Path
from hashlib import sha256
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch
from docx import Document

from core import commercial_notes
from tests.test_commercial_word import sample


class CommercialNotesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        folder = self.root / "commerciale" / "calabria"
        folder.mkdir(parents=True)
        self.target = folder / "SCALEA-sintetico.docx"
        self.original = sample()
        self.target.write_bytes(self.original)
        self.root_patch = patch.object(commercial_notes, "library_workspace_root", return_value=self.root)
        self.root_patch.start()
        self.addCleanup(self.root_patch.stop)

    def test_find_multiple_clients_in_one_file_and_select_specific_block(self):
        result = commercial_notes.find_customers("TEST")
        self.assertEqual(len(result["matches"]), 2)
        self.assertEqual({m["heading"] for m in result["matches"]},
                         {"LIDO TEST UNO", "HOTEL TEST DUE"})
        self.assertTrue(all(m["path"] == "commerciale/calabria/SCALEA-sintetico.docx"
                            for m in result["matches"]))

    def test_append_only_and_private_original_backup(self):
        match = commercial_notes.find_customers("LIDO TEST")[ "matches" ][0]
        row = {
            "id": uuid4(), "document_path": match["path"],
            "document_sha256": match["sha256"],
            "customer_heading": match["heading"],
            "customer_start": match["start"], "customer_stop": match["stop"],
            "note_date": "24/10/2026", "note_text": "Confermato richiamo cliente",
        }
        backup = commercial_notes._publish(row)
        before = [p.text for p in Document(BytesIO(self.original)).paragraphs]
        after = [p.text for p in Document(self.target).paragraphs]
        self.assertEqual(after[:4], before[:4])
        self.assertEqual(after[4:6], ["24/10/2026", "Confermato richiamo cliente"])
        self.assertEqual(after[6:], before[4:])
        self.assertEqual((self.root / backup).read_bytes(), self.original)
        self.assertTrue(backup.startswith(".cora-commercial-backups/"))

    def test_moved_or_modified_document_never_overwritten(self):
        match = commercial_notes.find_customers("HOTEL TEST")[ "matches" ][0]
        row = {
            "id": uuid4(), "document_path": match["path"],
            "document_sha256": match["sha256"],
            "customer_heading": match["heading"], "customer_start": match["start"],
            "customer_stop": match["stop"], "note_date": "24/10/2026",
            "note_text": "Nuova nota",
        }
        self.target.write_bytes(sample() + b"unsafe")
        changed = self.target.read_bytes()
        with self.assertRaises(ValueError):
            commercial_notes._publish(row)
        self.assertEqual(self.target.read_bytes(), changed)

    def test_reject_path_escape_or_hidden_backups(self):
        for path in ("../../etc/passwd.docx", ".cora-commercial-backups/x.docx", "/etc/a.docx"):
            with self.subTest(path=path), self.assertRaises(ValueError):
                commercial_notes._document(path)

    def test_must_not_silently_match_another_customer(self):
        results = commercial_notes.find_customers("HOTEL TEST")
        self.assertEqual(len(results["matches"]), 1)
        self.assertEqual(results["matches"][0]["heading"], "HOTEL TEST DUE")


if __name__ == "__main__":
    unittest.main()
