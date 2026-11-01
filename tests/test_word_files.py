from io import BytesIO
from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import json

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from lxml import etree
from core.server_files import router
from core.ia_library import router as library_router
from core.word_files import W, NS


class WordFilesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.base = Path(self.temp.name)
        self.folder = self.base / 'files'; self.folder.mkdir()
        self.library = self.base / 'library'; self.library.mkdir()
        self.originals = self.base / 'originals'; self.originals.mkdir()
        self.env = patch.dict(os.environ, {'CORA_FILES_TOKEN':'owner','CORA_FILE_ROOTS':json.dumps([{'id':'files','path':str(self.folder),'writable':True},{'id':'ro','path':str(self.folder),'writable':False}]),'CORA_LIBRARY_ORIGINALS_ROOT':str(self.originals),'CORA_KNOWLEDGE_ROOT':str(self.library)})
        self.env.start()
        app = FastAPI(); app.include_router(router); app.include_router(library_router)
        self.client = TestClient(app,headers={'Authorization':'Bearer owner'})
        self.prefix = '/api/v1/server/files'
        doc = Document()
        doc.add_heading('LIDO MARIO',level=1)
        p = doc.add_paragraph(); p.add_run('Cliente: ').bold=True; p.add_run('Mario').italic=True
        doc.add_paragraph('Appunto iniziale\tcliente\nSeconda riga')
        table = doc.add_table(rows=1,cols=2); table.cell(0,0).text='Telefono'; table.cell(0,1).text='12345'
        doc.add_heading('HOTEL ROSSI',level=1); doc.add_paragraph('Non cambiare questi appunti')
        doc.sections[0].footer.paragraphs[0].text='Footer invariato'
        self.path=self.folder/'clienti.docx'; doc.save(self.path)
        self.original=self.path.read_bytes()

    def tearDown(self):
        self.client.close(); self.env.stop(); self.temp.cleanup()

    def get(self,root='files',path='clienti.docx',prefix=None):
        return self.client.get((prefix or self.prefix)+'/word',params={'root_id':root,'path':path})

    def save(self,revision,edits=None,notes=None,root='files',prefix=None):
        return self.client.post((prefix or self.prefix)+'/word',json={'root_id':root,'path':'clienti.docx','revision':revision,'edits':edits or {},'notes':notes or []})

    def test_reading_order_tables_formatting_and_no_source_mutation(self):
        response=self.get(); self.assertEqual(response.status_code,200,response.text)
        data=response.json()
        self.assertIn('no-store',response.headers['cache-control'])
        self.assertEqual([b['kind'] for b in data['blocks']],['paragraph','paragraph','paragraph','table','paragraph','paragraph'])
        self.assertEqual(data['paragraphs'][1]['text'],'Cliente: Mario')
        self.assertEqual(data['blocks'][3]['rows'][0]['cells'][0]['blocks'][0]['kind'], 'paragraph')
        self.assertTrue(data['paragraphs'][1]['runs'][0]['bold'])
        self.assertTrue(data['paragraphs'][1]['runs'][1]['italic'])
        self.assertEqual(self.path.read_bytes(),self.original)

    def test_edit_notes_keep_other_customer_runs_tables_footer_and_backup(self):
        data=self.get().json()
        result=self.save(data['revision'],{'p1':'Cliente: Mario Rossi','p4':'67890'},[{'after':'p2','text':'01/11/2026\nRichiamare il cliente.'}])
        self.assertEqual(result.status_code,200,result.text)
        saved=result.json()
        self.assertEqual((self.folder/saved['backupPath']).read_bytes(),self.original)
        doc=Document(self.path)
        self.assertIn('Cliente: Mario Rossi',[p.text for p in doc.paragraphs])
        self.assertIn('Richiamare il cliente.',[p.text for p in doc.paragraphs])
        self.assertEqual(doc.tables[0].cell(0,1).text,'67890')
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertTrue(doc.paragraphs[1].runs[1].italic)
        self.assertEqual(doc.paragraphs[-1].text,'Non cambiare questi appunti')
        with zipfile.ZipFile(BytesIO(self.original)) as before, zipfile.ZipFile(self.path) as after:
            self.assertEqual(before.namelist(),after.namelist())
            for name in before.namelist():
                if name!='word/document.xml': self.assertEqual(before.read(name),after.read(name),name)
            a=etree.fromstring(before.read('word/document.xml')); b=etree.fromstring(after.read('word/document.xml'))
            old=a.findall('.//w:p',NS)[-1]; new=b.findall('.//w:p',NS)[-1]
            self.assertEqual(etree.tostring(old),etree.tostring(new))
        self.assertEqual(self.get().json()['revision'],saved['revision'])

    def test_stale_revision_readonly_and_symlink_fail_without_change(self):
        revision=self.get().json()['revision']
        self.assertFalse(self.get(root='ro').json()['editable'])
        self.assertEqual(self.save(revision,{'p1':'new'},root='ro').status_code,403)
        self.assertEqual(self.save('0'*64,{'p1':'new'}).status_code,409)
        self.assertEqual(self.path.read_bytes(),self.original)
        link=self.folder/'link.docx'
        try: link.symlink_to(self.path)
        except OSError: return
        self.assertEqual(self.get(path='link.docx').status_code,403)
        self.assertEqual(self.get(path='../x.docx').status_code,400)
        self.assertEqual(self.client.get(self.prefix+'/word',params={'root_id':'files','path':'clienti.docx'},headers={'Authorization':''}).status_code,401)

    def test_malformed_legacy_invalid_edits_and_failed_save(self):
        (self.folder/'old.doc').write_bytes(b'legacy')
        self.assertEqual(self.get(path='old.doc').status_code,415)
        (self.folder/'broken.docx').write_bytes(b'not a zip')
        self.assertEqual(self.get(path='broken.docx').status_code,422)
        revision=self.get().json()['revision']
        self.assertEqual(self.save(revision,{'p999':'bad'}).status_code,422)
        self.assertEqual(self.save(revision,{'p1':'bad\x00text'}).status_code,422)
        with patch('core.file_editor.os.replace',side_effect=PermissionError):
            self.assertEqual(self.save(revision,{'p1':'new'}).status_code,403)
        self.assertEqual(self.path.read_bytes(),self.original)
        self.assertFalse(list(self.folder.glob('.cora-edit-*')))

    def test_library_edit_never_changes_server_original(self):
        imported=self.client.post('/api/v1/library/files/import',json={'source_root_id':'files','source_path':'clienti.docx','destination':'clienti.docx'})
        self.assertEqual(imported.status_code,201,imported.text)
        prefix='/api/v1/library/files'
        revision=self.get(root='library',prefix=prefix).json()['revision']
        self.assertEqual(self.save(revision,{'p1':'Copia modificata'},root='library',prefix=prefix).status_code,200)
        self.assertEqual(self.path.read_bytes(),self.original)
        self.assertNotEqual((self.library/'clienti.docx').read_bytes(),self.original)

    def test_empty_word_accepts_notes_and_protected_document_is_readonly(self):
        doc=Document(); doc.save(self.path)
        revision=self.get().json()['revision']
        result=self.save(revision,notes=[{'after':'end','text':'Primo appunto\nSeconda riga'}])
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual([p.text for p in Document(self.path).paragraphs],['Primo appunto','Seconda riga'])
        doc=Document(self.path)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        protection=OxmlElement('w:documentProtection'); protection.set(qn('w:enforcement'),'1')
        doc.settings.element.append(protection); doc.save(self.path)
        opened=self.get().json(); self.assertFalse(opened['editable'])
        self.assertEqual(self.save(opened['revision'],{'p0':'blocked'}).status_code,403)

    def test_drawings_breaks_and_unedited_paragraph_xml_are_preserved(self):
        from docx.enum.text import WD_BREAK
        doc=Document(self.path)
        import base64
        picture = BytesIO(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aPioAAAAASUVORK5CYII='))
        doc.add_picture(picture)
        doc.add_paragraph('Prima pagina').add_run().add_break(WD_BREAK.PAGE)
        doc.save(self.path)
        data=self.get().json()
        self.assertFalse(data['paragraphs'][-1]['editable'])
        old=self.path.read_bytes()
        self.assertEqual(self.save(data['revision'],{'p1':'Cliente: Luigi'}).status_code,200)
        with zipfile.ZipFile(BytesIO(old)) as before, zipfile.ZipFile(self.path) as after:
            a=etree.fromstring(before.read('word/document.xml')); b=etree.fromstring(after.read('word/document.xml'))
            self.assertEqual(etree.tostring(a.findall('.//w:p',NS)[-1]),etree.tostring(b.findall('.//w:p',NS)[-1]))
            images = [name for name in before.namelist() if name.startswith('word/media/')]
            self.assertTrue(images)
            for name in images: self.assertEqual(before.read(name), after.read(name))
