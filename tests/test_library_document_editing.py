import io
import unittest
from docx import Document
from core.library_tools import _document_bytes, digest


class DocumentEditingUnitTests(unittest.TestCase):
    def test_text_append_keeps_original(self):
        self.assertEqual(_document_bytes('.md', 'secondo', original=b'primo').decode(),
                         'primo\nsecondo')
        self.assertEqual(_document_bytes('.txt', 'secondo', original=b'primo\n').decode(),
                         'primo\nsecondo')

    def test_word_append_preserves_paragraphs_and_tables(self):
        doc = Document()
        doc.add_heading('Cliente Rossi', level=1)
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Ordine 123'
        stream = io.BytesIO()
        doc.save(stream)
        updated = Document(io.BytesIO(_document_bytes('.docx', 'Nuovo appunto', original=stream.getvalue())))
        self.assertEqual(updated.paragraphs[0].text, 'Cliente Rossi')
        self.assertEqual(updated.paragraphs[-1].text, 'Nuovo appunto')
        self.assertEqual(updated.tables[0].cell(0, 0).text, 'Ordine 123')

    def test_full_document_read_in_chunks(self):
        from core import library_tools
        from unittest.mock import patch
        with patch.object(library_tools, 'source_file') as source:
            import tempfile
            from pathlib import Path
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / 'notes.md'
                path.write_text('ABCDEFGH', encoding='utf-8')
                source.return_value = path
                # Underlying document tool function can be tested without Ollama.
                from core.governance import executables
                implementation = next(item['fn'] for item in executables.values()
                                      if item['tool'].name == 'library_read_document_chunk')
                with patch.object(library_tools, 'operation') as op:
                    from contextlib import nullcontext
                    op.return_value = nullcontext()
                    a = implementation('notes.md', 0, 3)
                    b = implementation('notes.md', a['next_offset'], 8, a['sha256'])
                self.assertEqual(a['content'] + b['content'], 'ABCDEFGH')
                self.assertFalse(b['has_more'])

    def test_only_editable_extensions(self):
        for suffix in ('.pdf', '.xlsx', '.py'):
            with self.assertRaises(ValueError):
                _document_bytes(suffix, 'data')

    def test_checksum_changes_with_content(self):
        self.assertNotEqual(digest(b'old'), digest(b'new'))


if __name__ == '__main__':
    unittest.main()
