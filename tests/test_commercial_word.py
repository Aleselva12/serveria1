"""Conservative commercial Word regression tests. All examples are synthetic."""
import unittest
from hashlib import sha256
from io import BytesIO
from docx import Document
from core.commercial_word import parse, search, propose_append


def sample():
    document = Document()
    for line in (
        "LIDO TEST UNO", "Mario", "21/10/2026", "Vecchia annotazione da preservare",
        "---------------------------------------------------------",
        "HOTEL TEST DUE", "Lucia", "21/10/2026", "Seconda annotazione",
        "---------------------------------------------------------",
    ):
        document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class CommercialWordTests(unittest.TestCase):
    def test_two_customer_blocks(self):
        blocks = parse(sample())
        self.assertEqual([item.heading for item in blocks], ["LIDO TEST UNO", "HOTEL TEST DUE"])
        self.assertEqual(len(search(sample(), "hotel test")), 1)

    def test_append_only_to_selected_customer(self):
        source = sample()
        blocks = parse(source)
        changed = propose_append(source, blocks[0], "22/10/2026", "Richiesta preventivo")
        old = [p.text for p in Document(BytesIO(source)).paragraphs]
        new = [p.text for p in Document(BytesIO(changed)).paragraphs]
        self.assertEqual(new[:4], old[:4])
        self.assertEqual(new[4:6], ["22/10/2026", "Richiesta preventivo"])
        self.assertEqual(new[6:], old[4:])
        self.assertNotEqual(sha256(source).hexdigest(), sha256(changed).hexdigest())

    def test_changed_document_rejected(self):
        source = sample()
        old_block = parse(source)[0]
        modified = propose_append(source, old_block, "22/10/2026", "Nota")
        with self.assertRaisesRegex(ValueError, "File cambiato"):
            propose_append(modified, old_block, "23/10/2026", "Altra nota")

    def test_reject_bad_input(self):
        source = sample()
        block = parse(source)[0]
        for date, note in [("22/10/26", "Test"), ("22/10/2026", ""), ("22/10/2026", "x" * 5001)]:
            with self.assertRaises(ValueError):
                propose_append(source, block, date, note)


if __name__ == "__main__":
    unittest.main()
