import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4
from core import library_tools as lib
from core.governance import execute_capability, executables, approved_action
from search_agent import search_tools as research


class LibraryWriteTests(unittest.TestCase):
    def make_symlink(self, link, target, **kwargs):
        try:
            link.symlink_to(target, **kwargs)
        except OSError as error:
            if os.name == "nt" and (getattr(error, "winerror", None) == 1314 or error.errno in {1, 13}):
                self.skipTest("Windows requires Developer Mode or symlink privileges for this test")
            raise

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)/'library'; self.root.mkdir()
        for target, value in [('search_agent.search_tools.KNOWLEDGE_ROOT',self.root),
                              ('core.permissions.policy_overrides',lambda: {}),
                              ('core.operation_journal.begin',lambda *args: None),
                              ('core.operation_journal.finish',lambda *args, **kwargs: None)]:
            p=patch(target,value);p.start();self.addCleanup(p.stop)
        (self.root/'note.txt').write_text('prima')
        self.hash=lib.digest(b'prima')

    def invoke(self, capability, payload, approve=False):
        name='supervisor.'+capability
        token=None
        if approve:
            key,entry=next((k,e) for k,e in executables.items() if e['contract'].id==name)
            normalized=entry['input_model'].model_validate(payload).model_dump()
            token=approved_action.set(('supervisor',frozenset(entry['actions']),key,normalized))
        try:return execute_capability(name,payload)
        finally:
            if token:approved_action.reset(token)

    def test_new_document_writes_require_confirmation(self):
        for capability, payload in (
            ('library_create_document', {'relative_path':'created.md','content':'Nuovo appunto'}),
            ('library_append_document', {'relative_path':'note.txt','expected_sha256':self.hash,'content':'Appunto aggiunto'}),
            ('library_replace_document_text', {'relative_path':'note.txt','expected_sha256':self.hash,'old_text':'prima','new_text':'aggiornato'}),
        ):
            with self.subTest(capability=capability), patch('core.governance.propose', return_value={'id': uuid4()}):
                self.assertEqual(self.invoke(capability,payload).status,'pending')
        self.assertFalse((self.root/'created.md').exists())
        self.assertEqual((self.root/'note.txt').read_text(),'prima')

    def test_create_append_and_restore_previous_content(self):
        made=self.invoke('library_create_document', {'relative_path':'created.md','content':'Prima riga'},True)
        self.assertEqual(made.status,'ok')
        before=(self.root/'created.md').read_bytes()
        changed=self.invoke('library_append_document', {'relative_path':'created.md','expected_sha256':lib.digest(before),'content':'Seconda riga'},True)
        self.assertEqual(changed.status,'ok')
        self.assertEqual((self.root/'created.md').read_text(),'Prima riga\\nSeconda riga')
        self.assertEqual((self.root/'.cora-trash'/changed.native['previous_version_trash_id']/'content').read_bytes(),before)
        self.assertEqual(self.invoke('library_append_document', {'relative_path':'created.md','expected_sha256':lib.digest(before),'content':'Ancora'},True).status,'error')

    def test_word_replace_retains_style_and_rejects_mixed_runs(self):
        from docx import Document
        path=self.root/'notes.docx'
        doc=Document(); p=doc.add_paragraph();r=p.add_run('Evidenza');r.bold=True
        doc.save(path)
        result=self.invoke('library_replace_document_text',{'relative_path':'notes.docx','expected_sha256':lib.digest(path.read_bytes()),'old_text':'Evidenza','new_text':'Confermato'},True)
        self.assertEqual(result.status,'ok')
        after=Document(path).paragraphs[0]
        self.assertEqual(after.text,'Confermato')
        self.assertTrue(after.runs[0].bold)
        doc=Document();p=doc.add_paragraph();p.add_run('Parte ');p.add_run('forte').bold=True
        doc.save(path)
        mixed=self.invoke('library_replace_document_text',{'relative_path':'notes.docx','expected_sha256':lib.digest(path.read_bytes()),'old_text':'Parte forte','new_text':'Rimpiazzo'},True)
        self.assertEqual(mixed.status,'error')
        self.assertEqual(Document(path).paragraphs[0].text,'Parte forte')

    def test_copy_preserves_source_and_occupied_destination(self):
        (self.root/'copy.txt').write_text('existing')
        result=self.invoke('library_copy',{'relative_path':'note.txt','destination':'copy.txt'})
        self.assertEqual(result.status,'ok')
        self.assertEqual((self.root/'copy (1).txt').read_text(),'prima')
        self.assertEqual((self.root/'copy.txt').read_text(),'existing')
        self.assertEqual((self.root/'note.txt').read_text(),'prima')
        self.assertNotEqual((self.root/'note.txt').stat().st_ino,(self.root/'copy (1).txt').stat().st_ino)

    def test_modification_and_trash_only_propose_before_approval(self):
        for tool,payload in [('library_replace_text',{'relative_path':'note.txt','expected_sha256':self.hash,'old_text':'prima','new_text':'dopo'}),('library_trash',{'relative_path':'note.txt','expected_sha256':self.hash})]:
            with patch('core.governance.propose',return_value={'id':uuid4()}) as propose:
                self.assertEqual(self.invoke(tool,payload).status,'pending')
                propose.assert_called_once()
            self.assertEqual((self.root/'note.txt').read_text(),'prima')
            self.assertFalse((self.root/'.cora-trash').exists())

    def test_approved_edit_preserves_previous_version_in_ui_trash_format(self):
        r=self.invoke('library_replace_text',{'relative_path':'note.txt','expected_sha256':self.hash,'old_text':'prima','new_text':'dopo'},True)
        self.assertEqual(r.status,'ok')
        self.assertEqual((self.root/'note.txt').read_text(),'dopo')
        slot=self.root/'.cora-trash'/r.native['previous_version_trash_id']
        self.assertEqual((slot/'content').read_text(),'prima')
        self.assertEqual(json.loads((slot/'metadata.json').read_text())['path'],'note.txt')

    def test_changed_file_rejects_approved_edit_and_trash(self):
        (self.root/'note.txt').write_text('changed')
        for tool,payload in [('library_replace_text',{'relative_path':'note.txt','expected_sha256':self.hash,'old_text':'prima','new_text':'dopo'}),('library_trash',{'relative_path':'note.txt','expected_sha256':self.hash})]:
            self.assertEqual(self.invoke(tool,payload,True).status,'error')
        self.assertEqual((self.root/'note.txt').read_text(),'changed')

    def test_approved_trash_keeps_content_restorable(self):
        result=self.invoke('library_trash',{'relative_path':'note.txt','expected_sha256':self.hash},True)
        self.assertEqual(result.status,'ok')
        self.assertFalse((self.root/'note.txt').exists())
        slot=self.root/'.cora-trash'/result.native['trash']['id']
        self.assertEqual((slot/'content').read_text(),'prima')

    def test_traversal_symlinks_and_directories_are_rejected(self):
        outside=self.root.parent/'outside.txt';outside.write_text('outside')
        self.make_symlink(self.root / 'link.txt', outside)
        for path in ('../outside.txt','link.txt','.'):
            self.assertEqual(self.invoke('library_copy',{'relative_path':path,'destination':'copy.txt'}).status,'error')
        self.assertEqual(self.invoke('library_copy',{'relative_path':'note.txt','destination':'../copy.txt'}).status,'error')
        self.assertFalse((self.root.parent/'copy.txt').exists())

    def test_existing_word_append_requires_confirmation(self):
        with patch('core.governance.propose',return_value={'id':uuid4()}) as propose:
            result=execute_capability('local_research_agent.append_word_document',{'relative_path':'note.docx','content':'added','expected_sha256':self.hash})
            self.assertEqual(result.status,'pending');propose.assert_called_once()

    def test_approved_word_append_preserves_previous_version_and_rejects_stale_hash(self):
        from docx import Document
        document=Document();document.add_paragraph('originale')
        path=self.root/'note.docx';document.save(path)
        before=path.read_bytes()
        key,entry=next((k,e) for k,e in executables.items() if e['contract'].id=='local_research_agent.append_word_document')
        payload=entry['input_model'].model_validate({'relative_path':'note.docx','content':'aggiunta','expected_sha256':lib.digest(before)}).model_dump()
        token=approved_action.set(('local_research_agent',frozenset(entry['actions']),key,payload))
        try:
            result=execute_capability('local_research_agent.append_word_document',payload)
            self.assertEqual(result.status,'ok')
            self.assertEqual([p.text for p in Document(path).paragraphs],['originale','aggiunta'])
            self.assertEqual(next((self.root/'.cora-trash').glob('*/content')).read_bytes(),before)
            self.assertEqual(execute_capability('local_research_agent.append_word_document',payload).status,'error')
        finally:approved_action.reset(token)

    def test_word_creation_cannot_overwrite_concurrent_destination(self):
        publish=lib.publish_bytes
        def raced(destination,data,**kwargs):
            destination.write_bytes(b'concurrent file')
            return publish(destination,data,**kwargs)
        with patch.object(lib,'publish_bytes',side_effect=raced):
            result=execute_capability('local_research_agent.create_word_document',{'title':'Test','content':'new','filename':'new.docx'})
        self.assertEqual(result.status,'error')
        self.assertEqual((self.root/'new.docx').read_bytes(),b'concurrent file')
