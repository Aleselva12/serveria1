"""Agent writes restricted to IA working copies, with version-bound approvals."""
import hashlib
import io
from docx import Document
import json
import os
from pathlib import Path
import tempfile
import uuid
from datetime import datetime, timezone

from core.governance import agent_tool
from core.file_paths import library_workspace_root
from core.server_files import operation, trash_area
from core.confined_paths import confined_path
from core.runtime_context import ensure_runtime_active


def source_file(relative_path):
    ensure_runtime_active()
    from search_agent import search_tools as research
    path = research._safe_path(relative_path)
    if not path.is_file():
        raise ValueError('Seleziona un file normale nella Libreria IA.')
    if path.stat().st_size > research.MAX_FILE_BYTES:
        raise ValueError('File oltre il limite di dimensione configurato.')
    return path


def digest(data):
    return hashlib.sha256(data).hexdigest()


def checked_file(relative_path, expected_sha256):
    path = source_file(relative_path)
    data = path.read_bytes()
    if digest(data) != expected_sha256:
        raise ValueError('File cambiato: rileggilo e richiedi una nuova approvazione.')
    return path, data


def publish_bytes(destination, data, *, replace=False):
    # Publish complete files. A concurrent destination is never overwritten by a copy.
    fd, name = tempfile.mkstemp(dir=destination.parent, prefix='.cora-staging-')
    temporary = Path(name)
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        if replace:
            os.replace(temporary, destination)
        else:
            os.link(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def archive_previous(path, data, *, move=False):
    area = trash_area(library_workspace_root())
    identifier = uuid.uuid4().hex
    slot = confined_path(area, area/identifier)
    slot.mkdir()
    metadata = {'id':identifier, 'path':path.relative_to(library_workspace_root()).as_posix(),
                'deletedAt':datetime.now(timezone.utc).isoformat(), 'reason':'agent-trash' if move else 'before-edit'}
    (slot/'metadata.json').write_text(json.dumps(metadata), encoding='utf-8')
    if move:
        path.rename(slot/'content')
    else:
        publish_bytes(slot/'content', data)
    return metadata


@agent_tool('supervisor', capability='library_file_info', actions=('library_read',), effect='read', retry='safe')
def library_file_info(relative_path: str) -> dict:
    """Legge dimensione e SHA256 del file nella Libreria IA prima di proporre una modifica o il cestino."""
    with operation():
        path = source_file(relative_path)
        data = path.read_bytes()
        return {'path':relative_path, 'sha256':digest(data), 'size_bytes':len(data)}


@agent_tool('supervisor', capability='library_copy', actions=('library_read', 'library_copy'), effect='write', retry='never')
def library_copy(relative_path: str, destination: str) -> dict:
    """Crea una copia nella sola Libreria IA. Non modifica l'origine; assegna un nome libero se già occupato."""
    from search_agent.search_tools import _safe_path
    with operation():
        ensure_runtime_active()
        source = source_file(relative_path)
        data = source.read_bytes()
        requested = _safe_path(destination)
        if not requested.parent.is_dir() or requested == library_workspace_root():
            raise ValueError('La cartella di destinazione deve già esistere.')
        for index in range(1000):
            target = requested if index == 0 else requested.with_name(f'{requested.stem} ({index}){requested.suffix}')
            target = _safe_path(target.relative_to(library_workspace_root()).as_posix())
            try:
                publish_bytes(target, data)
                return {'status':'ok', 'path':target.relative_to(library_workspace_root()).as_posix(), 'source':relative_path}
            except FileExistsError:
                continue
        raise ValueError('Troppi nomi occupati: scegli una destinazione diversa.')


@agent_tool('supervisor', capability='library_replace_text', actions=('library_read', 'library_modify'), effect='write', retry='never')
def library_replace_text(relative_path: str, expected_sha256: str, old_text: str, new_text: str) -> dict:
    """Propone la sostituzione di un singolo estratto testuale UTF-8. Mostra prima/dopo in Attività; esegue solo dopo conferma sulla stessa versione SHA256."""
    from search_agent.search_tools import TEXT_EXTENSIONS, MAX_FILE_BYTES
    with operation():
        ensure_runtime_active()
        path, data = checked_file(relative_path, expected_sha256)
        if path.suffix.lower() not in TEXT_EXTENSIONS:
            raise ValueError('Modifica puntuale disponibile solo per file testuali UTF-8; non PDF o Word.')
        text = data.decode('utf-8')
        if not old_text or text.count(old_text) != 1:
            raise ValueError('Il testo da sostituire deve essere presente una sola volta.')
        updated = text.replace(old_text, new_text, 1).encode('utf-8')
        if len(updated) > MAX_FILE_BYTES or updated == data:
            raise ValueError('Modifica vuota o oltre il limite di dimensione.')
        backup = archive_previous(path, data)
        publish_bytes(path, updated, replace=True)
        return {'status':'ok', 'path':relative_path, 'sha256':digest(updated), 'previous_version_trash_id':backup['id']}


@agent_tool('supervisor', capability='library_trash', actions=('library_read', 'library_trash'), effect='write', retry='never')
def library_trash(relative_path: str, expected_sha256: str) -> dict:
    """Propone lo spostamento di un file della Libreria IA nel cestino, solo dopo conferma sulla versione letta. Nessuna eliminazione definitiva."""
    with operation():
        ensure_runtime_active()
        path, data = checked_file(relative_path, expected_sha256)
        return {'status':'ok', 'trash':archive_previous(path, data, move=True)}




@agent_tool('supervisor', capability='library_read_document_chunk',
            actions=('library_read',), effect='read', retry='safe')
def library_read_document_chunk(relative_path: str, offset: int = 0, max_chars: int = 12000,
                                expected_sha256: str = '') -> dict:
    """Legge davvero il contenuto di TXT/MD/Word a blocchi. Restituisce testo,
    posizione successiva e SHA256: ripetere con next_offset finché has_more è falso."""
    from search_agent.search_tools import _read_docx
    with operation():
        path = source_file(relative_path)
        data = path.read_bytes()
        sha = digest(data)
        if expected_sha256 and sha != expected_sha256:
            raise ValueError('File modificato dalla prima lettura: ricomincia.')
        if path.suffix.lower() == '.docx':
            full_text = _read_docx(path)
        elif path.suffix.lower() in {'.txt', '.md'}:
            full_text = data.decode('utf-8')
        else:
            raise ValueError('Usa library_read per gli altri formati.')
        if not (0 <= offset <= len(full_text)) or not (1 <= max_chars <= 12000):
            raise ValueError('Intervallo di lettura non valido.')
        end = min(len(full_text), offset + max_chars)
        return {'status': 'ok', 'path': relative_path, 'sha256': sha,
                'content': full_text[offset:end], 'offset': offset, 'next_offset': end,
                'total_chars': len(full_text), 'has_more': end < len(full_text)}


def _document_bytes(suffix: str, text: str, *, original: bytes | None = None) -> bytes:
    """Text and DOCX editor; existing Word paragraphs, tables and styles survive append."""
    if suffix == '.docx':
        document = Document(io.BytesIO(original)) if original is not None else Document()
        for line in text.splitlines():
            document.add_paragraph(line)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()
    if suffix not in {'.txt', '.md'}:
        raise ValueError('Sono supportati soltanto .txt, .md e .docx.')
    if original is None:
        return text.encode('utf-8')
    previous = original.decode('utf-8')
    separator = '' if not previous or previous.endswith('\n') else '\n'
    return (previous + separator + text).encode('utf-8')


@agent_tool('supervisor', capability='library_create_document',
            actions=('library_modify',), effect='write', retry='never')
def library_create_document(relative_path: str, content: str) -> dict:
    """CREA un nuovo TXT, Markdown o DOCX nella Libreria IA dopo conferma in Attività.
    Non sovrascrive mai file esistenti."""
    from search_agent.search_tools import _safe_path, MAX_FILE_BYTES
    with operation():
        ensure_runtime_active()
        path = _safe_path(relative_path)
        if path == library_workspace_root() or not path.parent.is_dir() or path.is_symlink():
            raise ValueError('Destinazione non valida o cartella non esistente.')
        if not content.strip():
            raise ValueError('Il contenuto non può essere vuoto.')
        data = _document_bytes(path.suffix.lower(), content)
        if len(data) > MAX_FILE_BYTES:
            raise ValueError('Documento oltre il limite.')
        publish_bytes(path, data, replace=False)
        return {'status': 'ok', 'path': relative_path, 'sha256': digest(data),
                'size_bytes': len(data), 'created': True}


@agent_tool('supervisor', capability='library_append_document',
            actions=('library_read', 'library_modify'), effect='write', retry='never')
def library_append_document(relative_path: str, expected_sha256: str, content: str) -> dict:
    """AGGIUNGE appunti in fondo a TXT, Markdown o Word dopo approvazione.
    SHA256 obbligatorio: file cambiati non vengono modificati. Salva la versione
    precedente nel cestino prima di pubblicare la nuova."""
    from search_agent.search_tools import MAX_FILE_BYTES
    with operation():
        ensure_runtime_active()
        path, original = checked_file(relative_path, expected_sha256)
        if not content.strip():
            raise ValueError('Il contenuto non può essere vuoto.')
        updated = _document_bytes(path.suffix.lower(), content, original=original)
        if len(updated) > MAX_FILE_BYTES or updated == original:
            raise ValueError('Documento invariato o oltre il limite.')
        version = archive_previous(path, original)
        publish_bytes(path, updated, replace=True)
        return {'status': 'ok', 'path': relative_path, 'sha256': digest(updated),
                'previous_version_trash_id': version['id'], 'appended': True}


@agent_tool('supervisor', capability='library_replace_document_text',
            actions=('library_read', 'library_modify'), effect='write', retry='never')
def library_replace_document_text(relative_path: str, expected_sha256: str,
                                  old_text: str, new_text: str) -> dict:
    """Sostituisce un testo che compare una sola volta in TXT/MD o in un singolo
    paragrafo DOCX, conservando il resto del documento. Richiede conferma."""
    from search_agent.search_tools import MAX_FILE_BYTES
    with operation():
        ensure_runtime_active()
        path, original = checked_file(relative_path, expected_sha256)
        if not old_text or old_text == new_text:
            raise ValueError('Seleziona un testo da sostituire diverso dal nuovo.')
        if path.suffix.lower() == '.docx':
            doc = Document(io.BytesIO(original))
            # Avoid silent loss of mixed-run formatting: only exact paragraph replacements.
            paragraphs = list(doc.paragraphs)
            matches = [p for p in paragraphs if p.text == old_text]
            if len(matches) != 1:
                raise ValueError('Word: sostituzione ammessa solo per un paragrafo intero univoco.')
            # Reassigning Paragraph.text destroys run-level formatting.
            # Only replace when a single run can safely carry the edit.
            runs = matches[0].runs
            if len(runs) > 1:
                raise ValueError('Word con formattazione mista: modifica non sicura.')
            if runs:
                runs[0].text = new_text
            else:
                matches[0].add_run(new_text)
            output = io.BytesIO()
            doc.save(output)
            updated = output.getvalue()
        elif path.suffix.lower() in {'.txt', '.md'}:
            body = original.decode('utf-8')
            if body.count(old_text) != 1:
                raise ValueError('Il testo deve comparire una sola volta.')
            updated = body.replace(old_text, new_text, 1).encode('utf-8')
        else:
            raise ValueError('Formato non modificabile.')
        if len(updated) > MAX_FILE_BYTES:
            raise ValueError('Documento oltre il limite.')
        version = archive_previous(path, original)
        publish_bytes(path, updated, replace=True)
        return {'status': 'ok', 'path': relative_path, 'sha256': digest(updated),
                'previous_version_trash_id': version['id']}

LIBRARY_WRITE_TOOLS = [library_file_info, library_copy, library_replace_text, library_trash,
                       library_read_document_chunk,
                       library_create_document, library_append_document, library_replace_document_text]
