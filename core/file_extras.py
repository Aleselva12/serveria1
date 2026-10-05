"""Bounded previews/search, private expiring links and resumable owner uploads."""
import hashlib
import json
import os
import re
import math
import secrets
import shutil
import time
import uuid
from pathlib import Path

from fastapi import File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel, ConfigDict, Field
from core.confined_paths import confined_path

# Imported only when make_router finishes defining its basic routes.
TEXT_SUFFIXES = {'.txt', '.md', '.json', '.csv', '.tsv', '.py', '.ts', '.tsx', '.js', '.css', '.yaml', '.yml', '.toml', '.log', '.xml'}
PREVIEW_TYPES = {'.png': 'image/png', '.jpg': 'image/jpeg', '.jpeg': 'image/jpeg', '.gif': 'image/gif', '.webp': 'image/webp', '.pdf': 'application/pdf', '.mp3': 'audio/mpeg', '.wav': 'audio/wav', '.m4a': 'audio/mp4', '.ogg': 'audio/ogg', '.mp4': 'video/mp4', '.webm': 'video/webm'}

class UploadStart(BaseModel):
    root_id: str
    path: str = ''
    filename: str = Field(min_length=1, max_length=240)
    size: int = Field(ge=0)
    sha256: str = Field(pattern=r'^[0-9a-f]{64}$')

class UploadMetadata(UploadStart):
    model_config = ConfigDict(extra='forbid', strict=True)
    area: str
    created: float
    publishing: bool = False
    result: dict | None = None


class ShareStart(BaseModel):
    root_id: str
    path: str
    hours: int = Field(default=24, ge=1, le=168)


def fingerprint(path):
    s = path.stat()
    return f'{s.st_dev}:{s.st_ino}:{s.st_size}:{s.st_mtime_ns}'


def register_extras(router, root_provider, *, allow_upload):
    from core.server_files import operation, root_for, resolve, exists, vacant, parts, node, RESERVED, upload_to
    def select(root_id, write=False):
        return root_for(root_id, write, root_provider)

    @router.get('/preview')
    def preview(root_id: str, path: str):
        with operation():
            _, base = select(root_id)
            target = resolve(base, path)
            exists(target)
            if not target.is_file(): raise HTTPException(400, 'Seleziona un file normale.')
            suffix = target.suffix.lower()
            headers = {'X-Content-Type-Options': 'nosniff', 'Cache-Control': 'private, no-store',
                       'Content-Security-Policy': "sandbox; default-src 'none'"}
            if suffix in TEXT_SUFFIXES or suffix == '.docx':
                if suffix == '.docx':
                    if target.stat().st_size > 5_000_000: raise HTTPException(413, 'Documento troppo grande per l’anteprima.')
                    from docx import Document
                    try:
                        import zipfile
                        with zipfile.ZipFile(target) as archive:
                            if sum(i.file_size for i in archive.infolist()) > 20_000_000: raise ValueError('Expanded document too large')
                        document = Document(target)
                        text = '\n'.join(p.text for p in document.paragraphs)
                        text += '\n' + '\n'.join(' | '.join(c.text for c in row.cells) for table in document.tables for row in table.rows)
                    except Exception as error: raise HTTPException(422, 'Documento Word non leggibile.') from error
                    truncated = len(text) > 120000
                    text = text[:120000]
                else:
                    with target.open('rb') as reader: data = reader.read(120001)
                    truncated = len(data) > 120000
                    text = data[:120000].decode('utf-8', errors='replace')
                return JSONResponse({'kind': 'text', 'text': text, 'truncated': truncated}, headers=headers)
            if target.stat().st_size > 50_000_000: raise HTTPException(413, 'File oltre 50 MB: scaricalo per aprirlo. Anteprima non disponibile.')
            media = PREVIEW_TYPES.get(suffix)
            if not media: raise HTTPException(415, 'Anteprima non disponibile per questo formato. Puoi scaricare il file.')
            return FileResponse(target, media_type=media, filename=target.name, content_disposition_type='inline', headers=headers)

    @router.get('/search')
    def search(root_id: str, path: str = '', query: str = Query(min_length=1, max_length=200), offset: int = Query(0, ge=0, le=10000), limit: int = Query(100, ge=1, le=500)):
        with operation():
            root, base = select(root_id)
            folder = resolve(base, path)
            exists(folder)
            if not folder.is_dir(): raise HTTPException(400, 'Il percorso non è una cartella.')
            results, scanned, truncated = [], 0, False
            # Bound traversal even when there are no matches; never follow links or reserved areas.
            def visit(directory):
                nonlocal scanned, truncated
                for child in directory.iterdir():
                    if scanned >= 10000:
                        truncated = True
                        return
                    scanned += 1
                    if child.name.casefold() in RESERVED or child.is_symlink(): continue
                    if child.is_file() or child.is_dir():
                        if query.casefold() in child.name.casefold(): results.append(node(base, child, root['writable']))
                        if child.is_dir():
                            # Avoid unbounded recursion on a deep filesystem tree.
                            if len(child.relative_to(folder).parts) >= 30:
                                truncated = True
                            else: visit(child)
            visit(folder)
            results.sort(key=lambda n: (n['kind']!='folder', n['path'].casefold()))
            return {'rootId': root_id, 'path': path, 'parentPath': path.rpartition('/')[0] if path else None,
                    'writable': root['writable'], 'total': len(results), 'items': results[offset:offset+limit], 'truncated': truncated, 'scanned': scanned}

    @router.post('/shares', status_code=201)
    def create_share(body: ShareStart):
        from core.database import db_connection
        with operation():
            _, base = select(body.root_id)
            target = resolve(base, body.path)
            exists(target)
            if not target.is_file(): raise HTTPException(400, 'Condivisione disponibile per singoli file.')
            token = secrets.token_urlsafe(32)
            with db_connection() as conn:
                row = conn.execute('''INSERT INTO file_shares(token_hash,area,root_id,path,fingerprint,expires_at)
                    VALUES(%s,%s,%s,%s,%s,NOW() + %s * INTERVAL '1 hour') RETURNING id,expires_at''',
                    (hashlib.sha256(token.encode()).hexdigest(), router.prefix, body.root_id, body.path, fingerprint(target), body.hours)).fetchone()
            return {'id': str(row['id']), 'expiresAt': row['expires_at'].isoformat(), 'url': router.prefix + '/shared/' + token, 'requiresLogin': True}

    @router.get('/shares')
    def list_shares(root_id: str):
        from core.database import db_connection
        select(root_id)
        with db_connection() as conn:
            rows = conn.execute('SELECT id,path,expires_at,revoked FROM file_shares WHERE area=%s AND root_id=%s ORDER BY created_at DESC LIMIT 200', (router.prefix, root_id)).fetchall()
        return {'items': [{'id': str(r['id']), 'path': r['path'], 'expiresAt': r['expires_at'].isoformat(), 'revoked': r['revoked']} for r in rows]}

    @router.delete('/shares/{share_id}')
    def revoke_share(share_id: uuid.UUID):
        from core.database import db_connection
        with db_connection() as conn:
            row = conn.execute('UPDATE file_shares SET revoked=TRUE WHERE id=%s AND area=%s RETURNING id', (share_id, router.prefix)).fetchone()
        if not row: raise HTTPException(404, 'Condivisione non trovata.')
        return {'revoked': True}

    @router.get('/shared/{token}')
    def shared_file(token: str):
        from core.database import db_connection
        if len(token) > 100: raise HTTPException(404, 'Condivisione non disponibile.')
        with db_connection() as conn:
            row = conn.execute('SELECT * FROM file_shares WHERE token_hash=%s AND area=%s AND revoked=FALSE AND expires_at>NOW()', (hashlib.sha256(token.encode()).hexdigest(), router.prefix)).fetchone()
        if not row: raise HTTPException(404, 'Condivisione scaduta, revocata o non disponibile.')
        with operation():
            _, base = select(row['root_id'])
            target = resolve(base, row['path'])
            exists(target)
            if not target.is_file() or fingerprint(target) != row['fingerprint']:
                raise HTTPException(409, 'Il file condiviso è cambiato. Crea un nuovo collegamento.')
            return FileResponse(target, filename=target.name, media_type='application/octet-stream', headers={'Cache-Control': 'private, no-store'})

    def staging(base):
        area = confined_path(base, base / '.cora-staging')
        if area.is_symlink(): raise HTTPException(403, 'Area temporanea non valida.')
        area.mkdir(exist_ok=True)
        return area

    def session(root_id, ident):
        if not re.fullmatch(r'[0-9a-f]{32}', ident): raise HTTPException(400, 'Identificativo upload non valido.')
        root, base = select(root_id, True)
        area = staging(base)
        slot = confined_path(area, area / ('upload-' + ident))
        if slot.is_symlink() or (slot / 'meta.json').is_symlink() or (slot / 'content').is_symlink(): raise HTTPException(403, 'Upload non valido.')
        exists(slot / 'meta.json')
        try:
            meta = UploadMetadata.model_validate_json((slot / 'meta.json').read_text()).model_dump()
            if not math.isfinite(meta['created']): raise ValueError('Invalid timestamp')
            if len(parts(meta['filename'])) != 1: raise ValueError('Invalid filename')
            parts(meta['path'])
        except (ValueError, OSError) as error: raise HTTPException(409, 'Metadati upload non leggibili. Verifica la destinazione.') from error
        if meta['area'] != router.prefix or meta['root_id'] != root_id: raise HTTPException(404, 'Upload non trovato.')
        if meta['created'] + 86400 < time.time():
            shutil.rmtree(slot)
            raise HTTPException(410, 'Upload scaduto. Seleziona nuovamente il file.')
        return root, base, slot, meta

    @router.post('/uploads', status_code=201)
    def start_upload(body: UploadStart):
        with operation():
            root, base = select(body.root_id, True)
            if body.size > int(os.getenv('CORA_FILES_MAX_UPLOAD_BYTES', '1073741824')): raise HTTPException(413, 'File oltre il limite configurato.')
            if len(parts(body.filename)) != 1: raise HTTPException(400, 'Nome file non valido.')
            folder = resolve(base, body.path)
            exists(folder)
            if not folder.is_dir(): raise HTTPException(400, 'Destinazione non valida.')
            vacant(resolve(base, '/'.join(filter(None, [body.path, body.filename]))))
            area = staging(base)
            slots = list(area.glob('upload-*'))
            active_size, active_count = 0, 0
            for slot in slots:
                if slot.is_symlink() or not slot.is_dir(): continue
                try:
                    slot = confined_path(area, slot)
                    metadata = confined_path(slot, slot/'meta.json')
                    meta = UploadMetadata.model_validate_json(metadata.read_text()).model_dump()
                    if meta['created']+86400 < time.time(): shutil.rmtree(slot)
                    elif meta.get('result') is None:
                        active_size += meta['size']
                        active_count += 1
                except (OSError, ValueError, KeyError, HTTPException): continue
            if active_count >= 8: raise HTTPException(429, 'Massimo otto upload in sospeso per risorsa.')
            if active_size + body.size > 2 * int(os.getenv('CORA_FILES_MAX_UPLOAD_BYTES', '1073741824')): raise HTTPException(429, 'Troppi upload in sospeso. Completa o annulla quelli precedenti.')
            ident = uuid.uuid4().hex
            slot = confined_path(area, area / ('upload-' + ident))
            slot.mkdir()
            meta = {**body.model_dump(), 'area': router.prefix, 'created': time.time()}
            (slot/'meta.json').write_text(json.dumps(meta))
            (slot/'content').touch()
            return {'id': ident, 'offset': 0, 'size': body.size, 'sha256': body.sha256}

    @router.get('/uploads/{ident}')
    def upload_status(ident: str, root_id: str):
        with operation():
            _, _, slot, meta = session(root_id, ident)
            return {'id': ident, 'offset': meta['size'] if meta.get('result') is not None else (slot/'content').stat().st_size, 'size': meta['size'], 'sha256': meta['sha256']}

    @router.post('/uploads/{ident}/chunks')
    def upload_chunk(ident: str, root_id: str = Form(...), offset: int = Form(...), file: UploadFile = File(...)):
        try:
            data = file.file.read(8*1024*1024+1)
            if not data or len(data) > 8*1024*1024: raise HTTPException(413, 'Blocco upload non valido (massimo 8 MiB).')
            with operation():
                _, _, slot, meta = session(root_id, ident)
                if meta.get('result') is not None or meta.get('publishing'): raise HTTPException(409, 'Upload già completato o in pubblicazione.')
                current = (slot/'content').stat().st_size
                if offset != current: raise HTTPException(409, 'Offset cambiato. Riprendi dal punto confermato dal server.')
                if current+len(data) > meta['size']: raise HTTPException(413, 'Blocco oltre la dimensione dichiarata.')
                with (slot/'content').open('ab') as writer:
                    writer.write(data)
                    writer.flush()
                    os.fsync(writer.fileno())
                return {'id': ident, 'offset': current+len(data), 'size': meta['size'], 'sha256': meta['sha256']}
        finally: file.file.close()

    @router.post('/uploads/{ident}/complete')
    def complete_upload(ident: str, root_id: str):
        with operation():
            root, base, slot, meta = session(root_id, ident)
            # Keep the completed receipt to disambiguate a lost response; no repeated publication.
            if meta.get('result') is not None: return meta['result']
            if meta.get('publishing'): raise HTTPException(409, 'Esito della pubblicazione non confermato. Verifica la destinazione prima di un nuovo upload.')
            if (slot/'content').stat().st_size != meta['size']: raise HTTPException(409, 'Upload incompleto.')
            with (slot/'content').open('rb') as reader:
                if hashlib.file_digest(reader, 'sha256').hexdigest() != meta['sha256']: raise HTTPException(422, 'Il contenuto non corrisponde al file selezionato. Annulla questo upload.')
            vacant(resolve(base, '/'.join(filter(None, [meta['path'], meta['filename']]))))
            meta['publishing'] = True
            (slot/'meta.json').write_text(json.dumps(meta))
            with (slot/'content').open('rb') as reader:
                upload = UploadFile(filename=meta['filename'], file=reader)
                if allow_upload:
                    result = upload_to(root, base, meta['path'], upload)
                else:
                    from core.ia_library import upload_copy
                    result = upload_copy(meta['path'], upload)
            meta['result'] = result
            (slot/'meta.json').write_text(json.dumps(meta))
            (slot/'content').unlink(missing_ok=True)
            return result

    @router.delete('/uploads/{ident}')
    def abandon_upload(ident: str, root_id: str):
        with operation():
            _, _, slot, _ = session(root_id, ident)
            shutil.rmtree(slot)
            return {'cancelled': True}
